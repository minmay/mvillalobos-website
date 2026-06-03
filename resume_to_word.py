import yaml
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

def create_resume(yaml_path, docx_path):
    try:
        with open(yaml_path, 'r') as f:
            data = yaml.safe_load(f)
    except Exception as e:
        print(f"Error reading YAML file: {e}")
        return

    doc = Document()

    # Styling the document
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Calibri'
    font.size = Pt(11)

    # Header
    name_heading = doc.add_heading(data.get('name', 'Resume'), 0)
    name_heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    contact_para = doc.add_paragraph()
    contact_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    email = data.get('email', '')
    mobile = data.get('mobile', '')
    contact_text = f"{email} " + (f"| {mobile}" if mobile else "")
    contact_para.add_run(contact_text)

    # Introduction
    if data.get('introduction'):
        doc.add_heading('Professional Summary', level=1)
        doc.add_paragraph(data['introduction'])

    # Objective
    if data.get('objective'):
        doc.add_heading('Objective', level=1)
        for obj in data['objective']:
            doc.add_paragraph(obj, style='List Bullet')

    # Work History
    if data.get('history'):
        doc.add_heading('Professional Experience', level=1)
        for job in data['history']:
            # Company and Title
            p = doc.add_paragraph()
            run = p.add_run(f"{job.get('name', '')} - {job.get('title', '')}")
            run.bold = True
            
            # Dates and Location
            dates = f"{job.get('start', '')} - {job.get('end', '')}"
            loc = job.get('location', '')
            meta_text = f"{dates} " + (f"| {loc}" if loc else "")
            doc.add_paragraph(meta_text)
            
            # Description
            if job.get('description'):
                for desc_item in job['description']:
                    if isinstance(desc_item, dict):
                        if 'prologue' in desc_item:
                            doc.add_paragraph(desc_item['prologue'])
                        if 'accomplishments' in desc_item:
                            for acc in desc_item['accomplishments']:
                                doc.add_paragraph(acc, style='List Bullet')
                    elif isinstance(desc_item, str):
                        doc.add_paragraph(desc_item)

    # Education
    if data.get('education'):
        doc.add_heading('Education', level=1)
        for edu in data['education']:
            p = doc.add_paragraph()
            run = p.add_run(f"{edu.get('name', '')}")
            run.bold = True
            
            degree = edu.get('degree', '')
            dates = f"{edu.get('start', '')} - {edu.get('end', '')}"
            doc.add_paragraph(f"{degree} ({dates})")

    # Skills
    if data.get('skills'):
        doc.add_heading('Skills', level=1)
        skills_list = data['skills']
        if isinstance(skills_list, list):
            for skill in skills_list:
                doc.add_paragraph(skill, style='List Bullet')
        else:
            doc.add_paragraph(str(skills_list))

    # Presentations
    if data.get('presentations'):
        doc.add_heading('Presentations', level=1)
        for pres in data['presentations']:
            p = doc.add_paragraph(style='List Bullet')
            p.add_run(f"{pres.get('name', '')}").bold = True
            if pres.get('link'):
                p.add_run(f" - {pres['link']}")

    # Affiliations
    if data.get('affiliations'):
        doc.add_heading('Affiliations', level=1)
        for aff in data['affiliations']:
            p = doc.add_paragraph()
            run = p.add_run(f"{aff.get('name', '')} - {aff.get('title', '')}")
            run.bold = True
            
            dates = f"{aff.get('start', '')} - {aff.get('end', '')}"
            doc.add_paragraph(dates)
            
            if aff.get('description'):
                for desc in aff['description']:
                    doc.add_paragraph(desc)

    doc.save(docx_path)
    print(f"Successfully created resume at: {docx_path}")

if __name__ == "__main__":
    import sys
    if len(sys.argv) < 3:
        print("Usage: python convert_resume.py <input_yaml> <output_docx>")
    else:
        create_resume(sys.argv[1], sys.argv[2])
